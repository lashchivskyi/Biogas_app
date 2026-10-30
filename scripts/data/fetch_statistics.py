"""
Refreshes official market statistics. Read-only summary by default; --output PATH
writes a candidate, --write atomically updates production after mandatory feeds succeed.
Retains consumer bands, latest published period and dated ECB reference FX.

Sources (all public, no key):
- Eurostat nrg_pc_205  Electricity prices, non-household, band IC (500–1 999 MWh/a), excl. VAT & recoverable taxes, EUR/kWh
- Eurostat nrg_pc_203  Natural gas prices, non-household, band I3 (10 000–99 999 GJ/a), excl. VAT & recoverable taxes, EUR/kWh
- Eurostat lc_lci_lev  Hourly labour cost, industry/construction/services (B–S excl. O), EUR/h
- Eurostat ert_bil_eur_a  EUR reference exchange rates, annual average (fallback only)
- US by state: EIA Electric Power Monthly 5.6.B, EIA industrial gas by state, BLS OEWS 51-8031 + ECEC (API), EIA STEO table 7a hubs (needs openpyxl)
- National EUR rates where the ECB has none: NBU (UAH), NBM (MDL), NBK (KZT), BCRA (ARS), NBS (RSD), SBV cross rate (VND, recorded by hand)
- Energy-Charts API (Fraunhofer ISE; ENTSO-E / SMARD data)  Day-ahead prices by bidding zone, EUR/MWh, last 12 months
"""
import argparse, concurrent.futures, json, re, time, urllib.request, urllib.error, datetime, statistics, math, calendar
from pathlib import Path
import xml.etree.ElementTree as ET

TARGET = Path(__file__).resolve().parents[2] / "src/countries/data/statistics.json"
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--output", type=Path)
parser.add_argument("--write", action="store_true")
parser.add_argument("--as-of", default=datetime.datetime.now(datetime.timezone.utc).date().isoformat())
args = parser.parse_args()
as_of = datetime.date.fromisoformat(args.as_of)
if args.write and args.output:
    parser.error("Choose --output or --write")
if args.output and args.output.resolve() == TARGET.resolve():
    parser.error("Production update requires --write")

EU = ["PL", "CZ", "SK", "HU", "RO", "MD", "BG", "RS", "EL", "TR", "LT", "LV", "EE", "DK", "SE", "FI", "DE", "AT", "IT", "FR", "ES", "NL"]
ISO = {"EL": "GR"}
CURRENCIES = ["PLN", "CZK", "HUF", "RON", "MDL", "RSD", "TRY", "DKK", "SEK", "GBP", "USD", "CAD", "MXN", "ARS", "BRL", "INR", "THB", "VND", "KZT", "UAH"]
BZN = {"PL": "PL", "CZ": "CZ", "SK": "SK", "HU": "HU", "RO": "RO", "BG": "BG", "RS": "RS", "GR": "GR", "LT": "LT", "LV": "LV", "EE": "EE",
       "DK": "DK1", "SE": "SE4", "FI": "FI", "DE": "DE-LU", "AT": "AT", "IT": "IT-North", "FR": "FR", "ES": "ES", "NL": "NL"}


def get(url, tries=2):
    for i in range(tries):
        try:
            with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "methalytics-data/2.0"}), timeout=25) as r:
                return json.load(r)
        except (urllib.error.URLError, TimeoutError) as e:
            if i == tries - 1:
                raise
            time.sleep(1)


def eurostat(dataset, query, key_dim="geo", bands=False):
    url = f"https://ec.europa.eu/eurostat/api/dissemination/statistics/1.0/data/{dataset}?format=JSON&lang=en&{query}"
    d = get(url)
    if d.get("updated", "")[:10] > args.as_of:
        raise ValueError(f"{dataset}: publication is later than the requested valuation date; a historical snapshot is required")
    dims = d["id"]
    size = d["size"]
    idx = {k: {i: c for c, i in d["dimension"][k]["category"]["index"].items()} for k in dims}
    out = {}
    for flat, v in d["value"].items():
        if v == 0:
            continue  # Eurostat reports 0 for bands with no consumers (e.g. ≥ 150 GWh in the Baltics): not a price.
        if not isinstance(v, (int, float)) or not math.isfinite(v) or v < 0:
            raise ValueError(f"{dataset}: invalid numeric observation")
        flat = int(flat)
        coords = {}
        for k, s in zip(reversed(dims), reversed(size)):
            coords[k] = idx[k][flat % s]
            flat //= s
        period = coords["time"]
        if len(period) == 4:
            period_end = period + "-12-31"
        elif "-S" in period:
            period_end = period[:4] + ("-06-30" if period.endswith("S1") else "-12-31")
        elif len(period) == 7:
            period_end = period + f"-{calendar.monthrange(int(period[:4]), int(period[5:]))[1]:02d}"
        else:
            period_end = period
        if period_end > args.as_of:
            continue  # Never label an incomplete/future annual period as an observed full-year average.
        group = coords[key_dim] + ("|" + coords["nrg_cons"] if bands else "")
        out.setdefault(group, {})[coords["time"]] = v
    if not out:
        raise ValueError(f"Empty mandatory feed: {dataset}")
    return out, d.get("updated"), url


def latest_mean(series, n):
    ks = sorted(series)[-n:]
    return round(statistics.mean(series[k] for k in ks), 4), ks


old = json.loads(TARGET.read_text(encoding="utf-8"))
res = {**old, "electricity": {}, "gas": {}, "labour": {}, "fx": dict(old["fx"]), "dayAhead": dict(old["dayAhead"]), "updated": dict(old.get("updated", {})), "refreshErrors": {}}
geos = "&".join(f"geo={g}" for g in EU)

el, upd, url = eurostat("nrg_pc_205", f"tax=X_VAT&currency=EUR&unit=KWH&sinceTimePeriod=2025-S1&{geos}", bands=True)
res["updated"]["nrg_pc_205"] = upd
for group, s in el.items():
    g, band = group.split("|")
    v, ks = latest_mean(s, 2)
    last = sorted(s)[-1]
    row = {"eurKwh": s[last], "periods": [last], "referenceMean": v, "referencePeriods": ks, "sourceUpdated": upd, "fetchedOn": args.as_of, "url": url}
    res["electricity"].setdefault(ISO.get(g, g), {"bands": {}})["bands"][band] = row
for row in res["electricity"].values():
    row.update(row["bands"]["MWH500-1999"])

gas, upd, url = eurostat("nrg_pc_203", f"tax=X_VAT&currency=EUR&unit=KWH&sinceTimePeriod=2025-S1&{geos}", bands=True)
res["updated"]["nrg_pc_203"] = upd
for group, s in gas.items():
    g, band = group.split("|")
    v, ks = latest_mean(s, 2)
    last = sorted(s)[-1]
    row = {"eurMwh": s[last] * 1000, "periods": [last], "referenceMean": v * 1000, "referencePeriods": ks, "sourceUpdated": upd, "fetchedOn": args.as_of, "url": url}
    res["gas"].setdefault(ISO.get(g, g), {"bands": {}})["bands"][band] = row
for row in res["gas"].values():
    row.update(row["bands"]["GJ10000-99999"])

lab, upd, url = eurostat("lc_lci_lev", f"unit=EUR&lcstruct=D1_D4_MD5&nace_r2=B-S_X_O&sinceTimePeriod=2023&{geos}")
res["updated"]["lc_lci_lev"] = upd
for g, s in lab.items():
    k = sorted(s)[-1]
    res["labour"][ISO.get(g, g)] = {"eurH": s[k], "year": k, "sourceUpdated": upd, "fetchedOn": args.as_of, "url": url}

fx, upd, url = eurostat("ert_bil_eur_a", "statinfo=AVG&sinceTimePeriod=2024", key_dim="currency")
res["updated"]["ert_bil_eur_a"] = upd
for c in CURRENCIES:
    if c in fx:
        k = sorted(fx[c])[-1]
        res["fx"][c] = {"perEur": fx[c][k], "year": k, "sourceUpdated": upd, "fetchedOn": args.as_of, "url": url}

fx_url = "https://www.ecb.europa.eu/stats/eurofxref/eurofxref-daily.xml"
try:
    with urllib.request.urlopen(fx_url, timeout=25) as response:
        tree = ET.fromstring(response.read())
    dated = next(node for node in tree.iter() if "time" in node.attrib)
    if dated.attrib["time"] > args.as_of:
        raise ValueError("Future ECB observation")
    res["fxReference"] = {"asOf": dated.attrib["time"], "fetchedOn": args.as_of, "url": fx_url, "rates": {node.attrib["currency"]: float(node.attrib["rate"]) for node in dated if "currency" in node.attrib}}
except Exception as exc:
    res["refreshErrors"]["fxReference"] = str(exc)[:200]
    # Verified HTTPS fallback: Eurostat republishes the ECB daily reference series.
    # Never disable TLS certificate verification to obtain a newer-looking date.
    daily, updated, daily_url = eurostat("ert_bil_eur_d", "statinfo=AVG&unit=NAC&sinceTimePeriod=" + str(as_of - datetime.timedelta(days=10)), key_dim="currency")
    periods = sorted({period for series in daily.values() for period in series if period <= args.as_of})
    latest = periods[-1]
    rates = {currency: series[latest] for currency, series in daily.items() if latest in series}
    res["fxReference"] = {"asOf": latest, "fetchedOn": args.as_of, "url": daily_url, "sourceUpdated": updated, "rates": rates}

def text(url):
    with urllib.request.urlopen(urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0 methalytics-data/2.0"}), timeout=25) as r:
        return r.read()


def back(fetch, days=7):
    """Latest quotation on or before the valuation date (central banks do not publish on holidays)."""
    for i in range(days):
        try:
            return fetch(as_of - datetime.timedelta(days=i))
        except (LookupError, ValueError, StopIteration, KeyError):
            continue
    raise LookupError("no quotation in the last week")


def nbu(day):
    url = "https://bank.gov.ua/NBUStatService/v1/statdirectory/exchange?valcode=EUR&date=" + day.strftime("%Y%m%d") + "&json"
    q = get(url)[0]
    return q["rate"], datetime.datetime.strptime(q["exchangedate"], "%d.%m.%Y").date().isoformat(), url, None


def nbm(day):
    url = "https://www.bnm.md/en/official_exchange_rates?get_xml=1&date=" + day.strftime("%d.%m.%Y")
    root = ET.fromstring(text(url))
    node = next(v for v in root.iter("Valute") if v.findtext("CharCode") == "EUR")
    return float(node.findtext("Value")) / float(node.findtext("Nominal")), datetime.datetime.strptime(root.attrib["Date"], "%d.%m.%Y").date().isoformat(), url, None


def nbk(day):
    url = "https://nationalbank.kz/rss/get_rates.cfm?fdate=" + day.strftime("%d.%m.%Y")
    root = ET.fromstring(text(url))
    node = next(i for i in root.iter("item") if i.findtext("title") == "EUR")
    return float(node.findtext("description")) / float(node.findtext("quant")), datetime.datetime.strptime(root.findtext("date"), "%d.%m.%Y").date().isoformat(), url, None


def bcra(day):
    url = "https://api.bcra.gob.ar/estadisticascambiarias/v1.0/Cotizaciones?fecha=" + day.isoformat()
    r = get(url)["results"]
    rate = next(x["tipoCotizacion"] for x in r["detalle"] if x["codigoMoneda"] == "EUR")
    if not rate:
        raise LookupError("no BCRA quotation")
    return rate, r["fecha"], url, None


def nbs(_day):
    # National Bank of Serbia official middle rate (www.nbs.rs blocks scripts; the NBS web-app host does not).
    url = "https://webappcenter.nbs.rs/ExchangeRateWebApp/ExchangeRate/CurrentMiddleRate"
    page = text(url).decode("utf-8", "ignore")
    rate = re.search(r">\s*EUR\s*<.*?>\s*978\s*<.*?>\s*1\s*<.*?>\s*([\d.]+,\d+)\s*<", page, flags=re.S)
    when = re.search(r"(\d{2})\.(\d{2})\.(\d{4})", page)
    if not rate or not when:
        raise LookupError("NBS middle rate not found")
    return float(rate.group(1).replace(".", "").replace(",", ".")), f"{when.group(3)}-{when.group(2)}-{when.group(1)}", url, None


res["fxNational"] = dict(old.get("fxNational", {}))
for currency, fetch, days in (("UAH", nbu, 7), ("MDL", nbm, 7), ("KZT", nbk, 7), ("ARS", bcra, 7), ("RSD", nbs, 1)):  # VND: SBV cross rate is recorded by hand (sbv.gov.vn rejects scripts)
    try:
        rate, quote_date, url, label = back(fetch, days)
        if quote_date > args.as_of or not rate or rate <= 0:
            raise ValueError("Invalid or future quotation")
        row = {"perEur": round(rate, 4), "asOf": quote_date, "fetchedOn": args.as_of, "url": url}
        if label:
            row["source"], row["confidence"] = label
        res["fxNational"][currency] = row
    except Exception as exc:
        res["refreshErrors"][f"fxNational.{currency}"] = str(exc)[:200]

end = as_of.replace(day=1)
start = end.replace(year=end.year - 1)
def spot(iso, bzn):
    time.sleep(6)  # Public API rate limit; keep the last successful row on failure.
    url = f"https://api.energy-charts.info/price?bzn={bzn}&start={start}&end={end}"
    d = get(url)
    p = [x for x in d.get("price", []) if x is not None]
    if not p:
        raise ValueError("Empty day-ahead series")
    return iso, {"eurMwh": statistics.mean(p), "bzn": bzn, "from": str(start), "to": str(end), "n": len(p), "fetchedOn": args.as_of, "url": url}

with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
    jobs = {pool.submit(spot, iso, bzn): iso for iso, bzn in BZN.items()}
    for job in concurrent.futures.as_completed(jobs):
        try:
            iso, row = job.result()
            res["dayAhead"][iso] = row
        except Exception as exc:
            res["refreshErrors"][f"dayAhead.{jobs[job]}"] = str(exc)[:200]

# ---- United States by state: prices and wages differ ×2–3 between states, so one national average misleads ----
US_STATES = {  # postal code (= GeoNames admin1 suffix): (name, FIPS, STEO wholesale hub)
    "AL": ("Alabama", "01", "ELWHU_SE"), "AK": ("Alaska", "02", None), "AZ": ("Arizona", "04", "ELWHU_SW"), "AR": ("Arkansas", "05", "ELWHU_MW"),
    "CA": ("California", "06", "ELWHU_CA"), "CO": ("Colorado", "08", "ELWHU_SW"), "CT": ("Connecticut", "09", "ELWHU_NE"), "DE": ("Delaware", "10", "ELWHU_PJ"),
    "DC": ("District of Columbia", "11", "ELWHU_PJ"), "FL": ("Florida", "12", "ELWHU_FL"), "GA": ("Georgia", "13", "ELWHU_SE"), "HI": ("Hawaii", "15", None),
    "ID": ("Idaho", "16", "ELWHU_NW"), "IL": ("Illinois", "17", "ELWHU_MW"), "IN": ("Indiana", "18", "ELWHU_MW"), "IA": ("Iowa", "19", "ELWHU_MW"),
    "KS": ("Kansas", "20", "ELWHU_SP"), "KY": ("Kentucky", "21", "ELWHU_SE"), "LA": ("Louisiana", "22", "ELWHU_MW"), "ME": ("Maine", "23", "ELWHU_NE"),
    "MD": ("Maryland", "24", "ELWHU_PJ"), "MA": ("Massachusetts", "25", "ELWHU_NE"), "MI": ("Michigan", "26", "ELWHU_MW"), "MN": ("Minnesota", "27", "ELWHU_MW"),
    "MS": ("Mississippi", "28", "ELWHU_MW"), "MO": ("Missouri", "29", "ELWHU_MW"), "MT": ("Montana", "30", "ELWHU_NW"), "NE": ("Nebraska", "31", "ELWHU_SP"),
    "NV": ("Nevada", "32", "ELWHU_SW"), "NH": ("New Hampshire", "33", "ELWHU_NE"), "NJ": ("New Jersey", "34", "ELWHU_PJ"), "NM": ("New Mexico", "35", "ELWHU_SW"),
    "NY": ("New York", "36", "ELWHU_NY"), "NC": ("North Carolina", "37", "ELWHU_SE"), "ND": ("North Dakota", "38", "ELWHU_MW"), "OH": ("Ohio", "39", "ELWHU_PJ"),
    "OK": ("Oklahoma", "40", "ELWHU_SP"), "OR": ("Oregon", "41", "ELWHU_NW"), "PA": ("Pennsylvania", "42", "ELWHU_PJ"), "RI": ("Rhode Island", "44", "ELWHU_NE"),
    "SC": ("South Carolina", "45", "ELWHU_SE"), "SD": ("South Dakota", "46", "ELWHU_SP"), "TN": ("Tennessee", "47", "ELWHU_SE"), "TX": ("Texas", "48", "ELWHU_TX"),
    "UT": ("Utah", "49", "ELWHU_SW"), "VT": ("Vermont", "50", "ELWHU_NE"), "VA": ("Virginia", "51", "ELWHU_PJ"), "WA": ("Washington", "53", "ELWHU_NW"),
    "WV": ("West Virginia", "54", "ELWHU_PJ"), "WI": ("Wisconsin", "55", "ELWHU_MW"), "WY": ("Wyoming", "56", "ELWHU_SW"),
}


def html_rows(raw):
    for tr in re.findall(r"<tr[^>]*>(.*?)</tr>", raw, flags=re.S):
        cells = [re.sub(r"<[^>]+>", " ", c).replace("&nbsp;", " ").strip() for c in re.findall(r"<t[dh][^>]*>(.*?)</t[dh]>", tr, flags=re.S)]
        if cells:
            yield [re.sub(r"\s+", " ", c) for c in cells]


def num(cell):
    try:
        return float(cell.replace(",", ""))
    except ValueError:
        return None


def us_electricity():
    url = "https://www.eia.gov/electricity/monthly/epm_table_grapher.php?t=epmt_5_6_b"
    raw = text(url).decode("utf-8", "ignore")
    title = re.search(r"Year-to-Date through (\w+) (\d{4})", raw)
    month = datetime.datetime.strptime(title.group(1), "%B").month
    year = int(title.group(2))
    names = {v[0]: k for k, v in US_STATES.items()}
    values = {}
    for cells in html_rows(raw):
        key = "US" if cells[0] == "U.S. Total" else ("US." + names[cells[0]] if cells[0] in names else None)
        if key and len(cells) >= 7 and num(cells[3]) and num(cells[5]):
            values[key] = {"commercial": round(num(cells[3]) / 100, 5), "industrial": round(num(cells[5]) / 100, 5)}  # columns: res, res-1, com, com-1, ind, ind-1, ...
    if len(values) < 50:
        raise ValueError("EIA table 5.6.B: too few states parsed")
    end = f"{year}-{month:02d}-{calendar.monthrange(year, month)[1]:02d}"
    return {"period": f"{year}-01/{year}-{month:02d}", "periodEnd": end, "url": url, "unit": "USD/kWh", "values": values}


def us_gas():
    heat_url = "https://www.eia.gov/dnav/ng/ng_cons_heat_a_EPG0_VGTH_btucf_a.htm"
    heat = {}
    raw = text(heat_url).decode("utf-8", "ignore")
    # Stub cells sit in nested tables: read "name … value cells" up to the history link.
    for name, block in re.findall(r'class="DataStub1">(?:<B>)?([^<]+)(?:<B>|</B>)?</td>(.*?)class="DataHist"', raw, flags=re.S):
        vals = [num(v) for v in re.findall(r'class="(?:DataB|Current2)">([^<]*)<', block) if num(v)]
        key = "US" if name.strip() == "U.S." else next(("US." + k for k, v in US_STATES.items() if v[0] == name.strip()), None)
        if key and vals:
            heat[key] = vals[-1] / 1000  # MMBtu per Mcf, latest year
    values, months = {}, {}
    for key, code in [("US", "us")] + [("US." + k, k.lower()) for k in US_STATES]:
        url = f"https://www.eia.gov/dnav/ng/hist/n3035{code}3m.htm"
        series = {}
        try:
            for cells in html_rows(text(url).decode("utf-8", "ignore")):
                if re.fullmatch(r"\d{4}", cells[0]) and len(cells) >= 13:
                    for m, c in enumerate(cells[1:13], start=1):
                        v = num(c)
                        if v:
                            series[f"{cells[0]}-{m:02d}"] = v
        except Exception:
            continue
        last = sorted(series)[-12:]
        if len(last) >= 9 and key in heat:  # ≥ 9 of the last 12 published months
            values[key] = round(sum(series[k] for k in last) / len(last) / (heat[key] * 0.293071), 3)
            months[key] = [last[0], last[-1], len(last)]
        time.sleep(0.3)
    if len(values) < 40:
        raise ValueError("EIA industrial gas: too few states")
    return {"unit": "USD/MWh HHV", "method": "Mean of the latest 12 published monthly industrial prices (USD/Mcf) ÷ (state heat content × 0.29307 MWh/MMBtu)",
            "url": "https://www.eia.gov/dnav/ng/ng_pri_sum_a_EPG0_PIN_DMcf_m.htm", "heatUrl": heat_url, "values": values, "months": months}


def us_operator():
    occ = {"US": "OEUN000000000000051803103", "ALL": "OEUN000000000000000000003"}
    occ.update({"US." + k: f"OEUS{v[1]}00000" + "000000" + "518031" + "03" for k, v in US_STATES.items()})  # area (7) + industry (6) + SOC + mean hourly wage
    ids = list(occ.values()) + ["CMU2010000000000D", "CMU2020000000000D"]
    found = {}
    for i in range(0, len(ids), 25):
        body = json.dumps({"seriesid": ids[i:i + 25], "startyear": str(as_of.year - 2), "endyear": str(as_of.year)}).encode()
        req = urllib.request.Request("https://api.bls.gov/publicAPI/v2/timeseries/data/", data=body, headers={"Content-Type": "application/json", "User-Agent": "methalytics-data/2.0"})
        with urllib.request.urlopen(req, timeout=40) as r:
            d = json.load(r)
        if d.get("status") != "REQUEST_SUCCEEDED":
            raise ValueError("BLS API: " + str(d.get("message"))[:150])
        for s in d["Results"]["series"]:
            if s["data"]:
                found[s["seriesID"]] = s["data"][0]
    total, wages = found["CMU2010000000000D"], found["CMU2020000000000D"]
    wage = {k: float(found[v]["value"]) for k, v in occ.items() if v in found and found[v]["value"] not in ("-", "")}
    year = found[occ["US"]]["year"]
    return {"occupation": "51-8031 Water and wastewater treatment plant and system operators", "year": year, "url": "https://www.bls.gov/oes/",
            "meanHourlyWage": wage, "ecec": {"period": f"{total['year']}-{total['period']}", "totalUsdH": float(total["value"]), "wagesUsdH": float(wages["value"]), "url": "https://www.bls.gov/ecec/"}}


def us_wholesale():
    import io
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("openpyxl is required for the EIA STEO workbook (pip install openpyxl)") from exc
    url = "https://www.eia.gov/outlooks/steo/xls/STEO_m.xlsx"
    wb = openpyxl.load_workbook(io.BytesIO(text(url)), read_only=True, data_only=True)
    dates = [r for r in wb["Dates"].iter_rows(max_row=6, values_only=True)]
    forecast = dates[1][3]  # modelling completion date; history ends the month before
    first = int(str(dates[4][3])[:4])
    end = (forecast.replace(day=1) - datetime.timedelta(days=1)).date()
    cols = [2 + (y - first) * 12 + (m - 1) for y, m in [((end.year * 12 + end.month - 1 - i) // 12, (end.year * 12 + end.month - 1 - i) % 12 + 1) for i in range(12)]]
    hubs = {}
    for r in wb["7atab"].iter_rows(values_only=True):
        if isinstance(r[0], str) and r[0].startswith("ELWHU_"):
            vals = [r[c] for c in cols if isinstance(r[c], (int, float))]
            if len(vals) == 12:
                hubs[r[0]] = {"name": r[1], "usdMwh": round(sum(vals) / 12, 2)}
    if len(hubs) < 10:
        raise ValueError("STEO table 7a: hubs missing")
    start = datetime.date(end.year - (1 if end.month < 12 else 0), end.month % 12 + 1, 1)
    return {"from": str(start), "to": str(end), "url": url, "table": "STEO table 7a, wholesale electricity prices", "hubs": hubs,
            "states": {"US." + k: v[2] for k, v in US_STATES.items() if v[2]}}


us = dict(old.get("usStates", {}))
for part, fetch in (("electricity", us_electricity), ("gas", us_gas), ("operator", us_operator), ("wholesale", us_wholesale)):
    try:
        us[part] = {**fetch(), "fetchedOn": args.as_of}
    except Exception as exc:
        res["refreshErrors"][f"usStates.{part}"] = str(exc)[:200]
us["names"] = {"US." + k: v[0] for k, v in US_STATES.items()}
res["usStates"] = us

res["fetchedOn"] = args.as_of
for mandatory in ("electricity", "gas", "labour"):
    missing = set(old[mandatory]) - set(res[mandatory])
    if missing:
        raise ValueError(f"{mandatory}: refusing to drop previously covered countries: {sorted(missing)}")
out = TARGET if args.write else args.output
if out:
    temporary = out.with_suffix(out.suffix + ".tmp")
    temporary.write_text(json.dumps(res, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    temporary.replace(out)
print(json.dumps({"written": str(out) if out else None, "counts": {key: len(res[key]) for key in ("electricity", "gas", "labour", "fx", "dayAhead")}, "fxReferenceDate": res.get("fxReference", {}).get("asOf"), "errors": res["refreshErrors"]}))
